from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from dengue_forecast.config import DISTRICTS, load_district_aliases, normalize_district_name
from dengue_forecast.contracts import ContractError

POPULATION_URL = "https://www.statistics.gov.lk/Population/StaticalInformation/CPH2024/Population_Tables"
POPULATION_SOURCE = "Sri Lanka Department of Census and Statistics CPH 2024 Population Tables"
POPULATION_SOURCE_VERSION = "CPH2024; workbook A1; retrieval verified 2026-10-04"
POPULATION_REFERENCE_YEAR = 2024
NATIONAL_TOTAL = 21_781_800


def _workbook(path: Path | str):
    try:
        content = Path(path).read_bytes()
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - exact openpyxl exception varies by bad payload
        raise ContractError(f"Could not read census workbook: {exc}") from exc


def _extract_english_name(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    aliases = sorted(load_district_aliases().values(), key=lambda item: len(item[1]), reverse=True)
    for _, canonical_name in aliases:
        if canonical_name in value:
            return canonical_name
    ascii_chunks = re.findall(r"[A-Za-z][A-Za-z .'-]*", value)
    if not ascii_chunks:
        return None
    name = re.sub(r"\s+", " ", ascii_chunks[-1]).strip()
    return name or None


def load_census_population(path: Path | str) -> pd.DataFrame:
    """Parse official 2024 Census workbook district population table A1."""

    wb = _workbook(path)
    if "A1" not in wb.sheetnames:
        raise ContractError("Census workbook missing A1 sheet")
    ws = wb["A1"]

    district_header = str(ws.cell(8, 1).value or "")
    total_header = str(ws.cell(8, 4).value or "")
    if "District" not in district_header or "Total number of persons" not in total_header:
        raise ContractError("Census A1 header validation failed")

    national_name = _extract_english_name(ws.cell(16, 1).value)
    national_total = ws.cell(16, 4).value
    if national_name != "Sri Lanka" or int(national_total) != NATIONAL_TOTAL:
        raise ContractError(f"Census national total mismatch: {national_name}={national_total}")

    records: list[dict[str, object]] = []
    for row in range(17, ws.max_row + 1):
        name = _extract_english_name(ws.cell(row, 1).value)
        value = ws.cell(row, 4).value
        if not name or name == "Sri Lanka":
            continue
        if value is None:
            continue
        try:
            district_id, district_name = normalize_district_name(name)
        except ContractError:
            continue
        population = int(value)
        if population <= 0:
            raise ContractError(f"Invalid population for {district_name}: {population}")
        records.append(
            {
                "district_id": district_id,
                "district_name": district_name,
                "population_reference": population,
                "population_reference_year": POPULATION_REFERENCE_YEAR,
                "population_method": "observed_retrospective_census_2024_no_interpolation",
                "population_source": POPULATION_SOURCE,
                "population_source_version": POPULATION_SOURCE_VERSION,
            }
        )

    out = pd.DataFrame(records).drop_duplicates("district_id", keep=False)
    expected_ids = {district.district_id for district in DISTRICTS}
    actual_ids = set(out["district_id"])
    if actual_ids != expected_ids:
        raise ContractError(f"Census district IDs mismatch: {sorted(expected_ids ^ actual_ids)}")
    if int(out["population_reference"].sum()) != NATIONAL_TOTAL:
        raise ContractError("Census district populations do not sum to national total")
    return out.sort_values("district_id").reset_index(drop=True)
